"""
model_proposed_labels.py — the trained fused model RE-LABELS every clip and we
put its proposal in a SEPARATE sheet next to the current label, disagreements
first, so the user adjudicates. Uses the cached 201-clip features (instant).

Output: training/birthday_model_proposed.xlsx
Columns: thumb | clip_id | current_label | model_proposed | agree | YOUR_final | notes
  - leave YOUR_final blank  → keep current_label
  - fill YOUR_final         → your decision wins (could be the model's, or your own)
Then finalize merges YOUR_final (else current_label) → labels_master.csv → retrain.
"""
import os, sys, glob, collections
import numpy as np
sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))

_HERE   = os.path.dirname(__file__)
_CACHE  = os.path.join(_HERE, "_fused_cache.npz")
_FRAMES = os.path.join(_HERE, "..", "..", "tests", "datasets", "Birthday", "Frames")
_OUT    = os.path.join(_HERE, "birthday_model_proposed.xlsx")
_MIN    = 3

TAXONOMY = ["pre_preparation", "person_arrival", "surprise_celebration", "cake",
            "cake_person", "cake_with_candles", "candle_blowing", "cake_cutting",
            "cake_smashing", "cake_feeding", "birthday_gifting", "group_photo",
            "individual_people", "smiling_moments", "laughing_moments",
            "crying_moments", "hugging_moments", "dancing_moments",
            "gazing_moments", "not_birthday"]


def _fit(X, y, iters=700, lr=0.5, l2=1e-3):
    w = np.zeros(X.shape[1], np.float32); b = 0.
    for _ in range(iters):
        p = 1/(1+np.exp(-(X@w+b))); g = p-y
        w -= lr*(X.T@g/len(y)+l2*w); b -= lr*float(g.mean())
    return w, b


def main():
    z = np.load(_CACHE, allow_pickle=True)
    X = z["X"]; cids = list(z["cids"]); ys = [str(v) for v in z["ys"]]
    counts = collections.Counter(ys); keep = {c for c, n in counts.items() if n >= _MIN}
    classes = sorted(keep); ci = {c: i for i, c in enumerate(classes)}
    mask = np.array([l in keep for l in ys])
    mu, sd = X[mask].mean(0), X[mask].std(0) + 1e-6
    Xs = (X[mask] - mu) / sd
    yc = np.array([ci[l] for l, m in zip(ys, mask) if m])
    # OUT-OF-FOLD predictions — each clip labelled by a model that NEVER trained
    # on it (honest judgement, not memorisation). This surfaces the real
    # disagreements worth adjudicating. Thin classes (<MIN) get a full-fit guess.
    learn_cids = [cid for cid, m in zip(cids, mask) if m]
    rng = np.random.RandomState(0); order = rng.permutation(len(yc)); K = 5
    proposed = {}
    for k in range(K):
        te = order[k::K]; tr = np.setdiff1d(order, te)
        Wk = np.zeros((len(classes), Xs.shape[1]), np.float32); Bk = np.zeros(len(classes), np.float32)
        for c in range(len(classes)): Wk[c], Bk[c] = _fit(Xs[tr], (yc[tr] == c).astype(np.float32))
        pr = np.argmax(Xs[te] @ Wk.T + Bk, axis=1)
        for j, idx in enumerate(te): proposed[learn_cids[idx]] = classes[int(pr[j])]
    # full-fit fallback for clips in classes too thin to fold
    Wf = np.zeros((len(classes), Xs.shape[1]), np.float32); Bf = np.zeros(len(classes), np.float32)
    for c in range(len(classes)): Wf[c], Bf[c] = _fit(Xs, (yc == c).astype(np.float32))
    for cid, v in zip(cids, X):
        proposed.setdefault(cid, classes[int(np.argmax(((v - mu) / sd) @ Wf.T + Bf))])

    rows = list(zip(cids, ys, [proposed[c] for c in cids]))
    # disagreements first, then grouped by current label
    rows.sort(key=lambda r: (r[1] == r[2], r[1], r[0]))

    from openpyxl import Workbook
    from openpyxl.worksheet.datavalidation import DataValidation
    from openpyxl.drawing.image import Image as XLImage
    wb = Workbook(); ws = wb.active; ws.title = "model_vs_you"
    ws.append(["thumb", "clip_id", "current_label", "model_proposed", "agree", "YOUR_final", "notes"])
    ws.freeze_panes = "A2"; ws.column_dimensions["A"].width = 26
    for col in "BCDF": ws.column_dimensions[col].width = 22
    dv = DataValidation(type="list", formula1='"' + ",".join(TAXONOMY) + '"', allow_blank=True)
    ws.add_data_validation(dv)
    n_dis = 0
    for i, (cid, cur, prop) in enumerate(rows, start=2):
        agree = (cur == prop)
        if not agree: n_dis += 1
        ws.cell(i,2,cid); ws.cell(i,3,cur); ws.cell(i,4,prop); ws.cell(i,5,"OK" if agree else "DIFF")
        dv.add(ws.cell(i,6)); ws.row_dimensions[i].height = 90
        fs = sorted(glob.glob(os.path.join(_FRAMES, cid, "*.jpg")))
        if fs:
            try: im = XLImage(fs[len(fs)//2]); im.width=160; im.height=120; ws.add_image(im, f"A{i}")
            except Exception: pass
    wb.save(_OUT)
    print(f"wrote {_OUT}: {len(rows)} clips, {n_dis} where model DIFFERS from current (shown first)")


if __name__ == "__main__":
    main()
