"""
train_fused.py — CLIP scene+action embedding FUSED with the discrimination cues
the user's notes described, turned into numbers:

  pout (mouthPucker/Funnel) → candle_blowing      jawOpen → cake_feeding
  browInnerUp/surprise      → surprise_celebration smile  → smiling_moments
  frosting-on-face          → cake_smashing        flame count → lit candles
  face size/bottom/count    → cake_person vs group / at-the-cake

Per clip = mean CLIP(scene⊕action) ⊕ aggregated hand-features. Standardised
PER-FOLD (no leakage). OVR logistic regression, leave-clips-out CV. Writes a
review sheet and reports vs the CLIP-only 34% baseline.
"""
from __future__ import annotations
import os, sys, glob, collections
import numpy as np

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))
import cv2
from numbers_parser import Document
from api.globals import CLIP, FACE_PROVIDER
from models.emotion_detector import detect_emotions
from models.candle_phase import count_warm_flames
from models.cake_smash import _frosting_fraction

_HERE    = os.path.dirname(__file__)
_FRAMES  = os.path.join(_HERE, "..", "..", "tests", "datasets", "Birthday", "Frames")
_NUMBERS = os.path.join(_HERE, "..", "..", "tests", "datasets", "Frames", "Birthday_frames", "birthday_labels.numbers")
_OUTX    = os.path.join(_HERE, "birthday_labels_fused.xlsx")
_CACHE   = os.path.join(_HERE, "_fused_cache.npz")
_MIN_PER_CLASS = 3

HAND = ["pucker", "funnel", "jaw", "brow", "surprise", "smile", "frost",
        "hfrac", "bottom", "flame_max", "flame_mean", "fc_mean", "fc_max",
        "blow_pout_x_lit", "pout_x_lean"]


_MASTER = os.path.join(_HERE, "labels_master.csv")

def load_labels():
    # Prefer the master CSV (145 corrected + newly-ingested pre-labelled clips).
    if os.path.exists(_MASTER):
        import csv
        out = {}
        for row in csv.DictReader(open(_MASTER)):
            cid = (row.get("clip_id") or "").strip()
            lab = (row.get("label") or "").strip()
            if cid and lab:
                out[cid] = dict(my="", final=lab, notes="")
        if out:
            return out
    t = Document(_NUMBERS).sheets[0].tables[0]
    out = {}
    for r in list(t.rows(values_only=True))[1:]:
        cid, my, your, notes = r[1], r[5], r[6], r[8]
        if not cid: continue
        out[str(cid)] = dict(my=str(my or ""), final=(str(your or "").strip() or str(my or "")),
                             notes=str(notes or ""))
    return out


def _emb(img):
    sc = CLIP.score_frame(img, event="birthday")
    e = getattr(sc, "_image_emb", None)
    return np.asarray(e, np.float32) if e is not None else None


def frame_vec(fr):
    """(clip_embed 1024, hand-feature dict) for one frame."""
    H, W = fr.shape[:2]
    scene = _emb(fr)
    if scene is None: return None, None
    ff = FACE_PROVIDER.detect(fr)
    faces = ff.faces if getattr(ff, "available", False) else []
    if faces:
        big = max(faces, key=lambda f: f.box[2] * f.box[3])
        x, y, w, h = big.box; pad = int(0.45 * max(w, h))
        crop = fr[max(0, y-pad):min(H, y+h+pad), max(0, x-pad):min(W, x+w+pad)]
    else:
        big = None; crop = fr[H//4:3*H//4, W//4:3*W//4]
    action = _emb(crop) if crop.size else scene
    if action is None: action = scene
    clip_vec = np.concatenate([scene, action])

    flames = float(count_warm_flames(fr))
    hd = dict(pucker=0., funnel=0., jaw=0., brow=0., surprise=0., smile=0.,
              frost=0., hfrac=0., bottom=0., flame=flames, fc=float(len(faces)))
    if big is not None:
        bs = big.blendshapes or {}
        hd["pucker"] = bs.get("mouthPucker", 0.); hd["funnel"] = bs.get("mouthFunnel", 0.)
        hd["jaw"] = bs.get("jawOpen", 0.)
        hd["brow"] = max(bs.get("browInnerUp", 0.), bs.get("browOuterUpLeft", 0.),
                         bs.get("browOuterUpRight", 0.))
        hd["smile"] = 0.5 * (bs.get("mouthSmileLeft", 0.) + bs.get("mouthSmileRight", 0.))
        x, y, w, h = big.box; hd["hfrac"] = h / H; hd["bottom"] = (y + h) / H
        hd["frost"] = _frosting_fraction(fr, big.box, W, H)
        try: hd["surprise"] = detect_emotions([big]).per_emotion.get("surprise", 0.)
        except Exception: pass
    return clip_vec, hd


def clip_fused(cid):
    fs = sorted(glob.glob(os.path.join(_FRAMES, cid, "*.jpg")))
    clips, hds = [], []
    for fp in fs:
        fr = cv2.imread(fp)
        if fr is None: continue
        cv_, hd = frame_vec(fr)
        if cv_ is None: continue
        clips.append(cv_); hds.append(hd)
    if not clips: return None
    cmean = np.mean(clips, axis=0); cmean /= (np.linalg.norm(cmean) + 1e-8)
    def col(k): return np.array([h[k] for h in hds], np.float32)
    # Interaction terms a LINEAR model can't form itself: the candle-blowing
    # signal is a pout AND lit candles in the SAME frame (not separately) — and
    # a pout while the face leans close to the cake. These split candle_blowing
    # from cake_with_candles (candles present, but no pout).
    pout = np.maximum(col("pucker"), col("funnel"))
    lit  = (col("flame") >= 1).astype(np.float32)
    hand = np.array([
        col("pucker").max(), col("funnel").max(), col("jaw").max(), col("brow").max(),
        col("surprise").max(), col("smile").max(), col("frost").max(),
        col("hfrac").max(), col("bottom").mean(), col("flame").max(), col("flame").mean(),
        col("fc").mean(), col("fc").max(),
        float((pout * lit).max()),              # pout WHILE candles lit (same frame)
        float((pout * col("hfrac")).max()),     # pout while leaning close to the cake
        ], np.float32)
    return np.concatenate([cmean, hand]).astype(np.float32)


def _fit(X, y, iters=700, lr=0.5, l2=1e-3):
    w = np.zeros(X.shape[1], np.float32); b = 0.
    for _ in range(iters):
        p = 1/(1+np.exp(-(X@w+b))); g = p-y
        w -= lr*(X.T@g/len(y)+l2*w); b -= lr*float(g.mean())
    return w, b
def _ovr(X, y, C):
    W = np.zeros((C, X.shape[1]), np.float32); B = np.zeros(C, np.float32)
    for c in range(C): W[c], B[c] = _fit(X, (y == c).astype(np.float32))
    return W, B
def _std(tr, te):
    mu = tr.mean(0); sd = tr.std(0) + 1e-6
    return (tr-mu)/sd, (te-mu)/sd


def main():
    if not CLIP.available: print("CLIP off."); return
    CLIP.warmup()
    labels = load_labels()
    if os.path.exists(_CACHE):
        z = np.load(_CACHE, allow_pickle=True); X = z["X"]; cids = list(z["cids"]); ys = list(z["ys"])
    else:
        X, cids, ys = [], [], []
        for k, (cid, info) in enumerate(labels.items()):
            v = clip_fused(cid)
            if v is None: continue
            X.append(v); cids.append(cid); ys.append(info["final"])
            if (k+1) % 25 == 0: print(f"  embedded {k+1}/{len(labels)}")
        X = np.stack(X); np.savez(_CACHE, X=X, cids=np.array(cids), ys=np.array(ys))
    print(f"{len(cids)} clips, feature dim {X.shape[1]} (1024 CLIP + {len(HAND)} cues)")

    counts = collections.Counter(ys); keep = {c for c, n in counts.items() if n >= _MIN_PER_CLASS}
    mask = np.array([l in keep for l in ys]); classes = sorted(keep); ci = {c: i for i, c in enumerate(classes)}
    Xc = X[mask]; yc = np.array([ci[l] for l, m in zip(ys, mask) if m])
    rng = np.random.RandomState(0); order = rng.permutation(len(yc)); K = 5
    correct = 0; conf = np.zeros((len(classes), len(classes)), int)
    for k in range(K):
        te = order[k::K]; tr = np.setdiff1d(order, te)
        Xtr, Xte = _std(Xc[tr], Xc[te])
        W, B = _ovr(Xtr, yc[tr], len(classes))
        pred = np.argmax(Xte @ W.T + B, axis=1); correct += int((pred == yc[te]).sum())
        for a, b in zip(yc[te], pred): conf[a, b] += 1
    acc = correct / max(1, len(yc))
    print(f"\n=== FUSED CV accuracy (leave-clips-out, {len(classes)} classes) = {acc:.3f}  (CLIP-only was 0.343) ===")
    for i, c in enumerate(classes):
        tot = conf[i].sum(); print(f"  {c:22s} {conf[i,i]}/{tot}" + (f"  ({conf[i,i]/tot:.0%})" if tot else ""))

    # full-fit predictions for the review sheet
    Xs = (Xc - Xc.mean(0)) / (Xc.std(0) + 1e-6)
    W, B = _ovr(Xs, yc, len(classes))
    mu, sd = Xc.mean(0), Xc.std(0) + 1e-6
    preds = {cid: classes[int(np.argmax(((v-mu)/sd) @ W.T + B))] for cid, v in zip(cids, X)}

    from openpyxl import Workbook
    from openpyxl.drawing.image import Image as XLImage
    wb = Workbook(); ws = wb.active; ws.append(
        ["thumb", "clip_id", "my_old", "your_label", "final", "fused_pred", "match", "notes"])
    ws.freeze_panes = "A2"; ws.column_dimensions["A"].width = 26
    i = 2
    for cid, info in labels.items():
        p = preds.get(cid, "")
        ws.cell(i,2,cid); ws.cell(i,3,info["my"]); ws.cell(i,4,info.get("your") or "")
        ws.cell(i,5,info["final"]); ws.cell(i,6,p); ws.cell(i,7,"OK" if p==info["final"] else "x")
        ws.cell(i,8,info["notes"]); ws.row_dimensions[i].height = 90
        fs = sorted(glob.glob(os.path.join(_FRAMES, cid, "*.jpg")))
        if fs:
            try: im = XLImage(fs[len(fs)//2]); im.width=160; im.height=120; ws.add_image(im, f"A{i}")
            except Exception: pass
        i += 1
    wb.save(_OUTX)
    print(f"review → {_OUTX}")


if __name__ == "__main__":
    main()
