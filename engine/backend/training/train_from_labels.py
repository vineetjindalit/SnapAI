"""
training/train_from_labels.py — rebuild ONE event's CLIP centroids from the
human-corrected labels in Birthday_frames/labels_visual.xlsx.

Per-event isolation (hard requirement): writes ONLY
    models/event_centroids/<event>.npz
so retraining birthday can never shift wedding's accuracy. The shared global
clip_class_centroids.npz ("general" fallback) is untouched.

Robust to file format: the labelling file may be a real .xlsx (Excel /
LibreOffice) OR an Apple Numbers file saved with an .xlsx extension — both are
read transparently.

Reads the "LABEL HERE (runs)" sheet (one row per continuous block) and expands
each run's label to all its frames; the "frames (detail)" sheet overrides
individual frames. Writes birthday_labels.csv (the per-frame label file) and
then trains.

Usage:
    python -m training.train_from_labels --event birthday [--val 0.15] [--force]
"""
from __future__ import annotations

import argparse
import csv
import json
import os
import shutil
import sys
import tempfile
from collections import Counter, defaultdict
from pathlib import Path

import numpy as np

sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..")))

ROOT = Path(__file__).resolve().parent.parent.parent
CENTROID_DIR = Path(__file__).resolve().parent.parent / "models" / "event_centroids"


def _find_xlsx() -> Path | None:
    env = os.environ.get("SNAPPY_FRAMES_XLSX")
    if env and Path(env).exists():
        return Path(env)
    hits = sorted(ROOT.glob("**/Birthday_frames/labels_visual.xlsx"))
    return hits[0] if hits else None


def _frame_idx(frame_file) -> int:
    try:
        return int(str(frame_file).split("_")[1].split(".")[0])
    except Exception:
        return -1


def _load_tables(path: Path) -> dict:
    """Return {sheet_name: [rows]} from a real .xlsx OR an Apple Numbers file."""
    # Try real Excel first.
    try:
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True)
        return {ws.title: [tuple(r) for r in ws.iter_rows(min_row=1, values_only=True)]
                for ws in wb.worksheets}
    except Exception:
        pass
    # Fall back to Numbers (needs a .numbers extension for the parser).
    from numbers_parser import Document
    tmp = Path(tempfile.mkdtemp()) / "labels.numbers"
    shutil.copy(path, tmp)
    doc = Document(str(tmp))
    return {s.name: s.tables[0].rows(values_only=True) for s in doc.sheets}


def read_final_labels(path: Path):
    """Expand runs→frames + apply detail overrides. Returns (labels, n_edits)."""
    tables = _load_tables(path)
    runs = tables.get("LABEL HERE (runs)", [])
    det  = tables.get("frames (detail)", [])

    final: dict = {}
    n_edits = 0

    # Runs sheet: [image, video, frames"a–b", count, trange, model_pred, label]
    for r in runs[1:]:
        if not r or len(r) < 7 or r[1] is None:
            continue
        video, frames, mpred, label = r[1], r[2], r[5], r[6]
        if not label:
            continue
        if mpred and str(label) != str(mpred):
            n_edits += 1
        try:
            a, b = (int(x) for x in str(frames).replace("–", "-").replace("—", "-").split("-"))
        except Exception:
            continue
        for idx in range(a, b + 1):
            final[(str(video), idx)] = str(label)

    # Detail sheet overrides: [image, video, frame_file, t, model_pred, label, conf]
    for r in det[1:]:
        if not r or len(r) < 6 or r[1] is None:
            continue
        video, ff, mpred, label = r[1], r[2], r[4], r[5]
        if label and mpred and str(label) != str(mpred):
            final[(str(video), _frame_idx(ff))] = str(label)
            n_edits += 1

    # Merge the focused negatives sheet (separate file, .numbers or .xlsx) as
    # the FINAL override. Prefer .numbers if it's newer (the user edits there).
    cands = [Path(path).parent / f"negative_label.{ext}" for ext in ("numbers", "xlsx")]
    cands = [p for p in cands if p.exists()]
    if cands:
        neg_path = max(cands, key=lambda p: p.stat().st_mtime)
        ov, n_ov = _read_overrides(neg_path)
        if ov:
            final.update(ov)
            n_edits += n_ov
            print(f"Merged {n_ov} edits from {neg_path.name} ({len(ov)} frames)")

    return final, n_edits


def _read_overrides(neg_path: Path):
    """negative_label.xlsx → {(video, idx): label} for rows the user changed.
       Columns: [image, video, frames, why, current, label]."""
    tables = _load_tables(neg_path)
    rows = (tables.get("negatives to label")
            or next((v for k, v in tables.items() if k != "vocab"), []))
    out, n = {}, 0
    for r in rows[1:]:
        if not r or len(r) < 6 or r[1] is None:
            continue
        video, frames, cur, label = r[1], r[2], r[4], r[5]
        if not label or str(label) == str(cur):
            continue
        try:
            a, b = (int(x) for x in str(frames).replace("–", "-").replace("—", "-").split("-"))
        except Exception:
            continue
        for idx in range(a, b + 1):
            out[(str(video), idx)] = str(label)
        n += 1
    return out, n


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--event", default="birthday")
    ap.add_argument("--val", type=float, default=0.15)
    ap.add_argument("--min-per-class", type=int, default=4)
    ap.add_argument("--force", action="store_true")
    args = ap.parse_args()

    xlsx = _find_xlsx()
    if xlsx is None:
        print("labels_visual.xlsx not found."); return
    print(f"Reading {xlsx}")
    final, n_edits = read_final_labels(xlsx)
    final = {k: v for k, v in final.items() if v and v.lower() != "none"}
    print(f"{len(final)} labelled frames, {n_edits} human edits detected")
    if n_edits < 20 and not args.force:
        print("\n⚠️  Looks UNLABELLED (few edits vs model predictions). "
              "Correct labels and re-save, or pass --force."); return

    print("Final per-frame label distribution:")
    for k, v in Counter(final.values()).most_common():
        print(f"  {k:24} {v}")

    # Write the per-frame label file (the deliverable).
    frames_dir = xlsx.parent
    csv_out = frames_dir / "birthday_labels.csv"
    with open(csv_out, "w", newline="") as f:
        w = csv.writer(f); w.writerow(["video", "frame_file", "label"])
        for (video, idx), lab in sorted(final.items()):
            w.writerow([video, f"frame_{idx:04d}.jpg", lab])
    print(f"Wrote per-frame labels → {csv_out}")

    # Embed each labelled frame via CLIP and build per-class centroids.
    import cv2
    from models.clip_engine import CLIPEngine
    clip = CLIPEngine.get(); clip.warmup()
    by_class = defaultdict(list)
    missing = 0
    for (video, idx), label in final.items():
        fp = frames_dir / video / f"frame_{idx:04d}.jpg"
        img = cv2.imread(str(fp))
        if img is None:
            missing += 1; continue
        s = clip.score_frame(img)
        emb = getattr(s, "_image_emb", None)
        if emb is not None:
            by_class[label].append(np.asarray(emb, dtype=np.float32))
    if missing:
        print(f"  ({missing} frames could not be read — skipped)")

    classes, centroids, val_sets = [], [], {}
    rng = np.random.default_rng(0)
    for cls, embs in by_class.items():
        if len(embs) < args.min_per_class:
            print(f"  skip '{cls}' — only {len(embs)} examples (<{args.min_per_class})")
            continue
        E = np.stack(embs); idx = rng.permutation(len(E))
        n_val = max(1, int(len(E) * args.val)) if len(E) >= 8 else 0
        val_idx, tr_idx = idx[:n_val], idx[n_val:]
        c = E[tr_idx].mean(0); c /= (np.linalg.norm(c) + 1e-8)
        classes.append(cls); centroids.append(c.astype(np.float32))
        val_sets[cls] = (val_idx, E)

    if not classes:
        print("No class had enough examples — nothing to train."); return
    C = np.stack(centroids).astype(np.float32)

    # Holdout: nearest-centroid accuracy on held-out frames.
    hits = tot = 0; per_cls = {}
    for cls, (val_idx, E) in val_sets.items():
        ch = ct = 0
        for vi in val_idx:
            pred = classes[int(np.argmax(C @ E[vi]))]
            tot += 1; ct += 1; hits += int(pred == cls); ch += int(pred == cls)
        if ct: per_cls[cls] = f"{ch}/{ct}"
    acc = hits / tot if tot else 0.0

    CENTROID_DIR.mkdir(parents=True, exist_ok=True)
    out = CENTROID_DIR / f"{args.event}.npz"
    np.savez(out, classes=np.array(classes, dtype=object), centroids=C,
             metrics=json.dumps({"val_acc": acc, "n_frames": len(final),
                                 "n_edits": n_edits, "classes": classes}))
    print(f"\nTrained {len(classes)} classes. Holdout accuracy: {acc:.3f} ({hits}/{tot})")
    for cls in classes:
        print(f"  {cls:24} val {per_cls.get(cls,'-')}")
    print(f"\nSaved → {out}")
    print("Restart the server; ONLY birthday sessions use this brain.")


if __name__ == "__main__":
    main()
