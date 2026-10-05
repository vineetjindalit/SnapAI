"""
review_master_labels.py — build ONE review sheet for all 201 clips (labels_master.csv),
grouped by current label, with a thumbnail + YOUR_label dropdown, so the user can
verify EVERY clip (incl. the auto-pulled pre-sorted ones) and correct outliers.

Output: training/birthday_review_all.xlsx
After the user fills YOUR_label (only where wrong) → finalize_master.py merges it.
"""
import os, sys, glob, csv
sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))
from openpyxl import Workbook
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.drawing.image import Image as XLImage

_HERE   = os.path.dirname(__file__)
_MASTER = os.path.join(_HERE, "labels_master.csv")
_NUMBERS= os.path.join(_HERE, "..", "..", "tests", "datasets", "Frames", "Birthday_frames", "birthday_labels.numbers")
_FRAMES = os.path.join(_HERE, "..", "..", "tests", "datasets", "Birthday", "Frames")
_OUT    = os.path.join(_HERE, "birthday_review_all.xlsx")

TAXONOMY = ["pre_preparation", "person_arrival", "surprise_celebration", "cake",
            "cake_person", "cake_with_candles", "candle_blowing", "cake_cutting",
            "cake_smashing", "cake_feeding", "birthday_gifting", "group_photo",
            "individual_people", "smiling_moments", "laughing_moments",
            "crying_moments", "hugging_moments", "dancing_moments",
            "gazing_moments", "not_birthday"]


def _reviewed_set():
    """clip_ids that came from the .numbers you already corrected (vs the new
    pre-sorted batch) — so the sheet can flag which still need a look."""
    try:
        from numbers_parser import Document
        t = Document(_NUMBERS).sheets[0].tables[0]
        return {str(r[1]) for r in list(t.rows(values_only=True))[1:] if r[1]}
    except Exception:
        return set()


def main():
    rows = []
    with open(_MASTER) as f:
        for r in csv.DictReader(f):
            rows.append((r["clip_id"], r["label"]))
    reviewed = _reviewed_set()
    rows.sort(key=lambda x: (x[1], x[0]))                 # group by label

    wb = Workbook(); ws = wb.active; ws.title = "review_all"
    ws.append(["thumb", "clip_id", "current_label", "YOUR_label", "source", "notes"])
    ws.freeze_panes = "A2"; ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 34; ws.column_dimensions["C"].width = 20
    ws.column_dimensions["D"].width = 20; ws.column_dimensions["F"].width = 40
    dv = DataValidation(type="list", formula1='"' + ",".join(TAXONOMY) + '"', allow_blank=True)
    ws.add_data_validation(dv)
    for i, (cid, lbl) in enumerate(rows, start=2):
        ws.cell(i, 2, cid); ws.cell(i, 3, lbl)
        ws.cell(i, 5, "you-reviewed" if cid in reviewed else "NEW-pre-sorted")
        dv.add(ws.cell(i, 4)); ws.row_dimensions[i].height = 90
        fs = sorted(glob.glob(os.path.join(_FRAMES, cid, "*.jpg")))
        if fs:
            try:
                im = XLImage(fs[len(fs) // 2]); im.width = 160; im.height = 120
                ws.add_image(im, f"A{i}")
            except Exception:
                pass
    wb.save(_OUT)
    n_new = sum(1 for c, _ in rows if c not in reviewed)
    print(f"wrote {_OUT}: {len(rows)} clips ({n_new} NEW pre-sorted to verify), grouped by label")


if __name__ == "__main__":
    main()
