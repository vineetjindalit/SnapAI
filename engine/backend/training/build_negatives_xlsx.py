"""
training/build_negatives_xlsx.py — a small, image-embedded sheet for labelling
just the negative-candidate runs (from negative_candidates.csv).

Each row = one candidate run with its middle-frame thumbnail, its current label,
and a dropdown to set the correct one (usually `negative`, or `cake` for a cake
close-up). Saved as negative_label.xlsx. train_from_labels merges these as
OVERRIDES on top of your full labels, so you don't touch the big sheet.
"""
from __future__ import annotations

import csv
from pathlib import Path

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from training.build_label_xlsx import (_thumb, _hdr, _make_vocab_sheet,
                                       THUMB_W, THUMB_H)

ROOT = Path(__file__).resolve().parent.parent.parent


def _frames_dir() -> Path | None:
    hits = sorted(ROOT.glob("**/Birthday_frames/negative_candidates.csv"))
    return hits[0].parent if hits else None


def main() -> None:
    fdir = _frames_dir()
    if fdir is None:
        print("negative_candidates.csv not found — run suggest_negatives first.")
        return
    rows = list(csv.DictReader(open(fdir / "negative_candidates.csv")))
    print(f"Building negatives sheet for {len(rows)} candidate runs")

    wb = Workbook()
    ws = wb.active; ws.title = "negatives to label"
    vocab_ref = _make_vocab_sheet(wb)   # hidden vocab sheet + range ref

    _hdr(ws, ["image", "video", "frames", "why", "current",
              "label (edit) → set negative / cake"])
    ws.column_dimensions["A"].width = THUMB_W / 7.0 + 1
    for col, w in zip("BCDEF", (34, 12, 16, 18, 30)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "B2"

    dv = DataValidation(type="list", formula1=vocab_ref, allow_blank=True)
    ws.add_data_validation(dv)
    afont = Font(name="Arial"); center = Alignment("center", vertical="center")
    flag = PatternFill("solid", fgColor="FFF2CC")

    buffers = []
    for i, r in enumerate(rows, start=2):
        video, frames, cur = r["video"], r["frames"], (r["current_label"] or "")
        try:
            a, b = (int(x) for x in frames.replace("–", "-").replace("—", "-").split("-"))
        except Exception:
            a = b = 0
        mid = (a + b) // 2
        fp = fdir / video / f"frame_{mid:04d}.jpg"
        try:
            buf, w, h = _thumb(fp); buffers.append(buf)
            img = XLImage(buf); img.width = w; img.height = h
            ws.add_image(img, f"A{i}")
        except Exception:
            ws.cell(row=i, column=1, value="(missing)")
        ws.row_dimensions[i].height = THUMB_H * 0.75
        ws.cell(row=i, column=2, value=video).font = afont
        ws.cell(row=i, column=3, value=frames).font = afont
        ws.cell(row=i, column=4, value=r.get("why", "")).font = afont
        ws.cell(row=i, column=5, value=cur).font = afont
        lab = ws.cell(row=i, column=6, value=cur)   # pre-fill = current label
        lab.font = Font(name="Arial", bold=True); lab.alignment = center
        lab.fill = flag
        dv.add(lab)
        for c in (3, 4, 5, 6):
            ws.cell(row=i, column=c).alignment = center

    out = fdir / "negative_label.xlsx"
    wb.save(out)
    print(f"Saved {out}  ({out.stat().st_size/1e6:.1f} MB, {len(rows)} runs)")
    print("Set the last column to 'negative' (or 'cake' for cake close-ups), "
          "save, and I'll merge + retrain.")


if __name__ == "__main__":
    main()
