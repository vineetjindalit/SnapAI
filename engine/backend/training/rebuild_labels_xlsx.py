"""
Rebuild birthday_labels.xlsx so the YOUR_label dropdown == your full birthday
taxonomy (matching the auto-label vocabulary), reusing the already-extracted
frames + auto-labels. Fast: no CLIP, no re-extraction.
"""
import os, sys, glob
sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))
from openpyxl import load_workbook, Workbook
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.drawing.image import Image as XLImage

_HERE = os.path.dirname(__file__)
_XLSX = os.path.join(_HERE, "birthday_labels.xlsx")
_FRAMES = os.path.join(_HERE, "..", "..", "tests", "datasets", "clips", "real", "Birthday", "frames")

# Your full birthday taxonomy (+ not_birthday for off-event clips).
TAXONOMY = ["pre_preparation", "person_arrival", "surprise_celebration", "cake",
            "cake_person", "cake_with_candles", "candle_blowing", "cake_cutting",
            "cake_smashing", "cake_feeding", "birthday_gifting", "group_photo",
            "individual_people", "smiling_moments", "laughing_moments",
            "crying_moments", "hugging_moments", "dancing_moments",
            "gazing_moments", "not_birthday"]
# normalize a few auto-label synonyms into the taxonomy
SYN = {"party": "group_photo", "decorations": "pre_preparation",
       "portrait": "individual_people", "gifting": "birthday_gifting",
       "general_peak": "group_photo", "hug_moment": "hugging_moments",
       "smiling": "smiling_moments", "laughing": "laughing_moments"}


def _mid_thumb(cid):
    fs = sorted(glob.glob(os.path.join(_FRAMES, cid, "*.jpg")))
    return fs[len(fs) // 2] if fs else None


def main():
    old = load_workbook(_XLSX).active
    rows = []
    for r in old.iter_rows(min_row=2, values_only=True):
        # cols: thumb, clip_id, dur, fps, n_frames, my_label, YOUR, conf, notes
        clip_id, dur, fps, nfr, my = r[1], r[2], r[3], r[4], r[5]
        if not clip_id:
            continue
        my = SYN.get(str(my), str(my))
        rows.append((clip_id, dur, fps, nfr, my, r[7]))

    wb = Workbook(); ws = wb.active; ws.title = "labels"
    ws.append(["thumb", "clip_id", "duration_s", "fps", "n_frames",
               "my_label", "YOUR_label", "auto_conf", "notes"])
    ws.freeze_panes = "A2"
    dv = DataValidation(type="list", formula1='"' + ",".join(TAXONOMY) + '"', allow_blank=True)
    ws.add_data_validation(dv)
    ws.column_dimensions["A"].width = 26; ws.column_dimensions["B"].width = 34
    for i, (cid, dur, fps, nfr, my, conf) in enumerate(rows, start=2):
        ws.cell(i, 2, cid); ws.cell(i, 3, dur); ws.cell(i, 4, fps)
        ws.cell(i, 5, nfr); ws.cell(i, 6, my); ws.cell(i, 8, conf)
        dv.add(ws.cell(i, 7)); ws.row_dimensions[i].height = 90
        th = _mid_thumb(cid)
        if th and os.path.exists(th):
            try:
                im = XLImage(th); im.width = 160; im.height = 120
                ws.add_image(im, f"A{i}")
            except Exception:
                pass
    wb.save(_XLSX)
    print(f"rebuilt {_XLSX} with {len(rows)} rows; dropdown = your taxonomy ({len(TAXONOMY)} moments)")


if __name__ == "__main__":
    main()
