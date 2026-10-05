"""
finalize_master.py — merge the user's YOUR_final decisions from
birthday_model_proposed.(xlsx|numbers) back into labels_master.csv.
final = YOUR_final if filled, else current_label. Then re-run train_fused.py.
"""
import os, sys, csv, glob, collections
_HERE   = os.path.dirname(__file__)
_MASTER = os.path.join(_HERE, "labels_master.csv")
_BASE   = "birthday_model_proposed"


def read_review():
    cands = [os.path.join(_HERE, _BASE + ext) for ext in (".numbers", ".xlsx")]
    cands = [p for p in cands if os.path.exists(p)]
    if not cands:
        cands = sorted(glob.glob(os.path.join(_HERE, "..", "..", "**", _BASE + ".*"),
                                 recursive=True), key=os.path.getmtime, reverse=True)
    path = sorted(cands, key=os.path.getmtime, reverse=True)[0]
    rows = []
    if path.endswith(".numbers"):
        from numbers_parser import Document
        t = Document(path).sheets[0].tables[0]
        src = list(t.rows(values_only=True))[1:]
    else:
        from openpyxl import load_workbook
        src = list(load_workbook(path).active.iter_rows(min_row=2, values_only=True))
    for r in src:
        # thumb, clip_id, current, model_proposed, agree, YOUR_final, notes
        rows.append((r[1], r[2], r[5] if len(r) > 5 else None))
    return path, rows


def main():
    path, rows = read_review()
    final, n = {}, 0
    for cid, cur, yf in rows:
        if not cid: continue
        yf = str(yf or "").strip()
        if yf and yf != str(cur or ""): n += 1
        final[str(cid)] = yf or str(cur or "")
    with open(_MASTER, "w", newline="") as f:
        w = csv.writer(f); w.writerow(["clip_id", "label"])
        for cid, lbl in final.items(): w.writerow([cid, lbl])
    print(f"read {path}")
    print(f"merged {len(final)} clips, {n} decisions applied → {_MASTER}")
    print("distribution:", dict(collections.Counter(final.values()).most_common()))


if __name__ == "__main__":
    main()
